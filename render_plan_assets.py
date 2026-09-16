# -*- coding: utf-8 -*-
"""
Generador de Capturas Ejecutivas en Alta Resolución (PNG) para el Plan de Medición Digital
Produce:
1. screenshots/09_plan_medicion_objetivos_kpis.png
2. screenshots/10_plan_medicion_matriz_eventos.png
3. screenshots/11_plan_medicion_dev_spec_datalayer.png
4. screenshots/12_arquitectura_data_pipeline_end_to_end.png
"""

import os
import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
import fitz  # PyMuPDF
import matplotlib.pyplot as plt
import matplotlib.patches as patches

os.makedirs("screenshots", exist_ok=True)
os.makedirs("scratch", exist_ok=True)

# 1. RENDERIZAR DIAGRAMA DE ARQUITECTURA PIPELINE
def generate_architecture_diagram():
    fig, ax = plt.subplots(figsize=(15, 8.5), dpi=220)
    fig.patch.set_facecolor('#0B1120')
    ax.set_facecolor('#0B1120')
    ax.axis('off')

    # Header
    ax.text(0.5, 0.94, 'ARQUITECTURA DE DATOS & PIPELINE ANALÍTICO END-TO-END', 
            ha='center', va='center', color='#F8FAFC', fontsize=18, fontweight='bold', family='sans-serif')
    ax.text(0.5, 0.88, 'CV Ejecutivo Inteligente | Telemetría, Gobernanza, BigQuery Warehouse y Capa de BI', 
            ha='center', va='center', color='#38BDF8', fontsize=12, family='sans-serif')

    boxes = [
        {
            "step": "PASO 1: CAPTURA",
            "title": "CV Web Interactivo\n(JavaScript DataLayer)",
            "details": "• Eventos snake_case\n• Desacople UI vs Datos\n• Nomenclatura corporativa\n• Cumplimiento No PII",
            "x": 0.11, "color": "#0284C7", "border": "#38BDF8"
        },
        {
            "step": "PASO 2: TAG MANAGEMENT",
            "title": "Google Tag Manager\n(GTM-P2Z4TZ4Z)",
            "details": "• 8 Tags GA4 Event\n• 8 Custom Triggers\n• 17 DataLayer Variables\n• Inyección asíncrona",
            "x": 0.305, "color": "#4F46E5", "border": "#818CF8"
        },
        {
            "step": "PASO 3: ANALÍTICA DIGITAL",
            "title": "Google Analytics 4\n(G-NQC5PHY67R)",
            "details": "• Dimensiones custom\n• Métricas de engagement\n• Medición de scroll\n• Exportación streaming",
            "x": 0.50, "color": "#D97706", "border": "#FBBF24"
        },
        {
            "step": "PASO 4: DATA WAREHOUSE",
            "title": "Google Cloud BigQuery\n(Dataset: analytics_tic)",
            "details": "• Esquema particionado\n• Desanidado UNNEST()\n• Detección de bots y ruidos\n• Funnel SQL Analytics",
            "x": 0.695, "color": "#059669", "border": "#34D399"
        },
        {
            "step": "PASO 5: VISUALIZACIÓN BI",
            "title": "Looker Studio &\nExecutive Dashboard",
            "details": "• Vistas limpias reales\n• Tasa de conversión neta\n• Métricas de reclutadores\n• Insights de negocio",
            "x": 0.89, "color": "#E11D48", "border": "#FB7185"
        }
    ]

    y_box = 0.50
    w_box = 0.17
    h_box = 0.44

    for b in boxes:
        # Box background
        rect = patches.FancyBboxPatch((b["x"] - w_box/2, y_box - h_box/2), w_box, h_box,
                                      boxstyle='round,pad=0.015,rounding_size=0.03',
                                      facecolor='#1E293B', edgecolor=b["border"], linewidth=2.2)
        ax.add_patch(rect)
        
        # Step banner pill
        pill = patches.FancyBboxPatch((b["x"] - 0.07, y_box + h_box/2 - 0.05), 0.14, 0.04,
                                      boxstyle='round,pad=0.01,rounding_size=0.02',
                                      facecolor=b["color"], edgecolor='none')
        ax.add_patch(pill)
        ax.text(b["x"], y_box + h_box/2 - 0.03, b["step"], ha='center', va='center',
                color='#FFFFFF', fontsize=8, fontweight='bold', family='sans-serif')
        
        # Title
        ax.text(b["x"], y_box + 0.08, b["title"], ha='center', va='center',
                color='#F1F5F9', fontsize=10.5, fontweight='bold', family='sans-serif', multialignment='center')
        
        # Details
        ax.text(b["x"] - 0.07, y_box - 0.10, b["details"], ha='left', va='center',
                color='#94A3B8', fontsize=9, family='sans-serif', linespacing=1.6)

    # Arrows between boxes
    for i in range(len(boxes) - 1):
        x_from = boxes[i]["x"] + w_box/2 + 0.005
        x_to = boxes[i+1]["x"] - w_box/2 - 0.005
        ax.annotate('', xy=(x_to, y_box), xytext=(x_from, y_box),
                    arrowprops=dict(arrowstyle='->', lw=2.8, color='#38BDF8', mutation_scale=16))

    # Real-Time Telemetry Sub-pipeline
    rect_tg = patches.FancyBboxPatch((0.20, 0.12), 0.60, 0.11,
                                     boxstyle='round,pad=0.01,rounding_size=0.02',
                                     facecolor='#0F172A', edgecolor='#0EA5E9', linewidth=1.5, linestyle='--')
    ax.add_patch(rect_tg)
    ax.text(0.50, 0.19, '⚡ CANAL REACTIVO COMPLEMENTARIO: Telemetría Instantánea vía Telegram Bot API',
            ha='center', va='center', color='#38BDF8', fontsize=9.5, fontweight='bold', family='sans-serif')
    ax.text(0.50, 0.145, 'Notificaciones push en tiempo real ante eventos de alto valor (Descarga de CV en PDF o Clic de Contacto) protegiendo la privacidad.',
            ha='center', va='center', color='#94A3B8', fontsize=8.5, family='sans-serif')

    # Footer
    ax.text(0.5, 0.04, 'Rubén David Barrios Bello | Lic. en Banca y Finanzas | Data Analyst & Business Intelligence Specialist',
            ha='center', va='center', color='#64748B', fontsize=9, family='sans-serif')

    out_path = "screenshots/12_arquitectura_data_pipeline_end_to_end.png"
    plt.tight_layout()
    plt.savefig(out_path, facecolor='#0B1120', edgecolor='none')
    plt.close()
    print(f"Generated: {out_path}")

# Helper para renderizar tablas estilizadas desde ReportLab a PNG
def render_table_pdf_to_png(pdf_path, png_path, title, subtitle, headers, rows, col_widths, page_size=landscape(letter)):
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=page_size,
        leftMargin=30,
        rightMargin=30,
        topMargin=26,
        bottomMargin=26
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'HeaderTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#0F172A')
    )
    sub_style = ParagraphStyle(
        'HeaderSub',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.5,
        leading=12,
        textColor=colors.HexColor('#0284C7')
    )
    cell_header = ParagraphStyle(
        'CellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
        alignment=1 # Center
    )
    cell_body = ParagraphStyle(
        'CellBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1E293B')
    )
    cell_bold = ParagraphStyle(
        'CellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#0F172A')
    )
    cell_code = ParagraphStyle(
        'CellCode',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7,
        leading=8.5,
        textColor=colors.HexColor('#0369A1')
    )

    story = [
        Paragraph(f"<b>{title}</b>", title_style),
        Spacer(1, 2),
        Paragraph(subtitle, sub_style),
        Spacer(1, 10)
    ]

    table_data = [[Paragraph(f"<b>{h}</b>", cell_header) for h in headers]]
    
    for r in rows:
        row_cells = []
        for i, val in enumerate(r):
            val_str = str(val or '').strip()
            if i == 0:
                p = Paragraph(val_str, cell_bold)
            elif '`' in val_str or '()' in val_str or '_' in val_str or '{' in val_str:
                clean_str = val_str.replace('<', '&lt;').replace('>', '&gt;')
                p = Paragraph(clean_str, cell_code)
            else:
                clean_str = val_str.replace('<', '&lt;').replace('>', '&gt;')
                p = Paragraph(clean_str, cell_body)
            row_cells.append(p)
        table_data.append(row_cells)

    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 0), (-1, 0), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 1), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 4),
    ]

    for idx in range(1, len(table_data)):
        bg = colors.HexColor('#F8FAFC') if idx % 2 == 1 else colors.white
        t_style.append(('BACKGROUND', (0, idx), (-1, idx), bg))

    t.setStyle(TableStyle(t_style))
    story.append(t)
    doc.build(story)

    # Convertir a PNG con PyMuPDF a 220 DPI
    pdf_doc = fitz.open(pdf_path)
    pix = pdf_doc[0].get_pixmap(dpi=220)
    pix.save(png_path)
    pdf_doc.close()
    print(f"Generated: {png_path} ({pix.width}x{pix.height})")

# 2. RENDERIZAR HOJA OBJETIVOS & KPIS
def generate_kpi_sheet():
    wb = openpyxl.load_workbook('Plan_de_Medicion_Digital_CV_Ejecutivo_Inteligente.xlsx', data_only=True)
    ws = wb['🎯 Objetivos & KPIs']
    
    headers = [ws.cell(3, c).value for c in range(1, 7)]
    rows = []
    for r in range(4, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 7)]
        if any(vals):
            rows.append(vals)

    col_widths = [60, 170, 200, 95, 135, 70]
    render_table_pdf_to_png(
        "scratch/kpis.pdf",
        "screenshots/09_plan_medicion_objetivos_kpis.png",
        "PLAN DE MEDICIÓN DIGITAL | 🎯 OBJETIVOS DE NEGOCIO, CONVERSIONES & KPIS",
        "Alineación Estratégica: Definición del Funnel de Empleabilidad y Métricas Clave de Desempeño",
        headers,
        rows,
        col_widths
    )

# 3. RENDERIZAR HOJA MATRIZ DE EVENTOS
def generate_events_matrix_sheet():
    wb = openpyxl.load_workbook('Plan_de_Medicion_Digital_CV_Ejecutivo_Inteligente.xlsx', data_only=True)
    ws = wb['⚡ Matriz de Eventos (GA4 & GTM)']
    
    headers = [ws.cell(3, c).value for c in range(1, 7)]
    rows = []
    # Seleccionamos los eventos representativos principales para que encaje perfecto en alta resolución
    for r in range(4, min(ws.max_row + 1, 19)):
        vals = [ws.cell(r, c).value for c in range(1, 7)]
        if any(vals):
            # Limpiamos texto para que luzca óptimo
            vals[5] = (vals[5] or '').replace('\n', ' • ')
            rows.append(vals)

    col_widths = [50, 90, 150, 115, 140, 185]
    render_table_pdf_to_png(
        "scratch/events.pdf",
        "screenshots/10_plan_medicion_matriz_eventos.png",
        "PLAN DE MEDICIÓN DIGITAL | ⚡ MATRIZ MAESTRA DE EVENTOS, VARIABLES & GTM",
        "Especificación Técnica de Eventos de Captura: Disparadores en Google Tag Manager y Parámetros GA4",
        headers,
        rows,
        col_widths
    )

# 4. RENDERIZAR HOJA DEV SPEC DATALAYER
def generate_dev_spec_sheet():
    wb = openpyxl.load_workbook('Plan_de_Medicion_Digital_CV_Ejecutivo_Inteligente.xlsx', data_only=True)
    ws = wb['💻 Dev Spec (DataLayer)']
    
    headers = [ws.cell(3, c).value for c in range(1, 5)]
    rows = []
    for r in range(4, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 5)]
        if any(vals):
            # Formatear snippets
            vals[2] = (vals[2] or '').replace('\n', '<br/>')
            rows.append(vals)

    col_widths = [120, 150, 270, 190]
    render_table_pdf_to_png(
        "scratch/devspec.pdf",
        "screenshots/11_plan_medicion_dev_spec_datalayer.png",
        "PLAN DE MEDICIÓN DIGITAL | 💻 ESPECIFICACIÓN TÉCNICA DEL DATALAYER (DEV SPEC)",
        "Contrato de Integración JavaScript para el Equipo de Desarrollo: Wrappers y Cargas Asíncronas",
        headers,
        rows,
        col_widths
    )

if __name__ == '__main__':
    print("Iniciando generación de activos visuales...")
    generate_architecture_diagram()
    generate_kpi_sheet()
    generate_events_matrix_sheet()
    generate_dev_spec_sheet()
    print("¡Todos los activos visuales fueron generados exitosamente!")
